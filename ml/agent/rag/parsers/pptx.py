"""PPTX parser implementation using python-pptx."""

import io
from typing import Union
from pptx import Presentation
from ml.agent.rag.parsers.base import BaseParser


class PptxParser(BaseParser):
    """Parser for PPTX documents."""

    async def parse(self, source: Union[str, bytes]) -> str:
        """
        Parse PPTX content and extract text.

        Parameters
        ----------
        source : Union[str, bytes]
            The PPTX content as bytes or file path.

        Returns
        -------
        str
            Extracted text from the PPTX.
        """
        if isinstance(source, bytes):
            pptx_file = io.BytesIO(source)
        else:
            pptx_file = source  # type: ignore

        prs = Presentation(pptx_file)
        text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
        return "\n".join(text)
